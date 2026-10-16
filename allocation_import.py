#!/usr/bin/env python3
"""Импорт аллокации сотрудников по проектам из Excel в нормализованный вид.

Исходный файл: на одном листе подряд идут месячные блоки (заголовок месяца,
шапка с названиями проектов, строки сотрудников с долями времени).
Результат: таблицы "сотрудник, проект, месяц, доля" в CSV и SQLite плюс отчет
report.md со всеми найденными проблемами и номерами строк исходного файла.

Пример:
    python allocation_import.py data/allocation.xlsx --out out
"""
from __future__ import annotations

import argparse
import ast
import contextlib
import csv
import difflib
import operator
import re
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# Столбцы листа (нумерация с 1): A тег, B ФИО, C должность, D дата найма,
# E..L доли по проектам.
COL_TAG, COL_NAME, COL_POSITION, COL_HIRE = 1, 2, 3, 4
FIRST_PROJECT_COL, LAST_PROJECT_COL = 5, 12
SHEET_WIDTH = 13

TOLERANCE = 0.005          # допустимое отклонение суммы долей от 1
DELIM = ";"                # разделитель CSV, чтобы файлы открывались в русском Excel
IGNORED_HEADERS = {"аллокация по проектам"}
DUPLICATE_TAG = "дубль"

MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
TITLE_RE = re.compile(r"^(?P<word>[а-я]+)[\s.]*(?P<year>\d{4}|\d{2})?(?P<fc>\s+прогноз)?$")
NOTE_IN_POSITION_RE = re.compile(r"ставк|убира|аутстафф|гпх|до \d")

ERROR, WARNING, INFO = "error", "warning", "info"


# --------------------------------------------------------------------------
# Модель данных
# --------------------------------------------------------------------------

@dataclass
class Issue:
    level: str
    month: str
    row: int
    message: str


@dataclass
class SourceRow:
    month: str
    forecast: bool
    row: int
    tag: str
    name: str
    position: str
    hire: date | None
    shares: dict[str, float]          # название проекта как в шапке -> доля


@dataclass
class Result:
    employee_months: list[dict] = field(default_factory=list)
    allocations: list[dict] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    raw_projects: dict[str, set[str]] = field(default_factory=dict)   # проект -> месяцы
    names: dict[str, str] = field(default_factory=dict)               # ключ -> ФИО
    months: list[str] = field(default_factory=list)
    raw_header: dict[str, list[str]] = field(default_factory=dict)          # месяц -> названия проектов в порядке столбцов
    project_positions: dict[str, list[str]] = field(default_factory=dict)   # месяц -> проекты аллокации в порядке столбцов

    def add(self, level: str, month: str, row: int, message: str) -> None:
        self.issues.append(Issue(level, month, row, message))


# --------------------------------------------------------------------------
# Разбор значений
# --------------------------------------------------------------------------

def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def key(value) -> str:
    return clean(value).casefold().replace("ё", "е")


_OPS = {ast.Add: operator.add, ast.Sub: operator.sub,
        ast.Mult: operator.mul, ast.Div: operator.truediv}


def _eval(node):
    if isinstance(node, ast.Expression):
        return _eval(node.body)
    if (isinstance(node, ast.Constant) and isinstance(node.value, (int, float))
            and not isinstance(node.value, bool)):
        return node.value
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        value = _eval(node.operand)
        return value if isinstance(node.op, ast.UAdd) else -value
    if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_eval(node.left), _eval(node.right))
    raise ValueError("формула не поддерживается")


def eval_formula(text: str) -> float:
    """Считает простую арифметику вида =1*0.5, если Excel не сохранил результат."""
    try:
        return float(_eval(ast.parse(text.lstrip("=").replace(",", "."), mode="eval")))
    except (SyntaxError, ZeroDivisionError) as exc:
        raise ValueError("формулу не удалось вычислить") from exc


def parse_share(value, formula=None) -> float | None:
    """Возвращает долю 0..1 или None для пустой ячейки. Ошибки: ValueError."""
    if value is None and isinstance(formula, str) and formula.startswith("="):
        value = formula
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"не число: {value!r}")
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        text = clean(value)
        if not text:
            return None
        try:
            if text.startswith("="):
                number = eval_formula(text)
            elif text.endswith("%"):
                number = float(text[:-1].replace(",", ".")) / 100
            else:
                number = float(text.replace(",", "."))
        except ValueError:
            raise ValueError(f"не число: {text!r}") from None
    else:
        raise ValueError(f"не число: {value!r}")
    if number < -1e-9 or number > 1 + 1e-9:
        raise ValueError(f"доля вне диапазона 0..1: {number:g}")
    return round(max(number, 0.0), 6)


def parse_title(value) -> tuple[int | None, int, bool] | None:
    """Заголовок месячного блока -> (год, месяц, прогноз) или None, если это не заголовок.

    В файле встречаются даты (26.03.2026 на деле означает "март 26") и текст:
    "май.26", "Июнь 2026", "Август 2026 ПРОГНОЗ", "Октябрь 26".
    """
    if isinstance(value, datetime):
        return value.year, value.month, False
    if not isinstance(value, str):
        return None
    match = TITLE_RE.match(clean(value).casefold())
    if not match or match["word"][:3] not in MONTHS:
        return None
    year = match["year"]
    if year is not None:
        year = int(year) + (2000 if len(year) == 2 else 0)
    return year, MONTHS[match["word"][:3]], bool(match["fc"])


# --------------------------------------------------------------------------
# Справочники (необязательные файлы настроек)
# --------------------------------------------------------------------------

def read_csv(path: Path, required: set[str]) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=DELIM)
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise SystemExit(f"{path}: нет столбцов {', '.join(sorted(missing))}")
        return [row for row in reader if any((v or "").strip() for v in row.values())]


def load_aliases(path: Path) -> dict[str, str]:
    """Файл с колонками alias;canonical: разные написания одного ФИО."""
    return {key(r["alias"]): clean(r["canonical"])
            for r in read_csv(path, {"alias", "canonical"})}


def load_project_map(path: Path) -> dict[str, list[tuple[str, float]]]:
    """Файл с колонками raw_name;project;weight.

    Одно название из шапки может превращаться в несколько проектов с весами
    (когда проект разделили), и несколько названий могут вести в один проект
    (когда проекты объединили или переименовали).
    """
    mapping: dict[str, list[tuple[str, float]]] = {}
    for r in read_csv(path, {"raw_name", "project"}):
        weight = clean(r.get("weight") or "1").replace(",", ".")
        try:
            mapping.setdefault(key(r["raw_name"]), []).append((clean(r["project"]), float(weight)))
        except ValueError:
            raise SystemExit(f"{path}: неверный вес {weight!r} у {r['raw_name']!r}") from None
    for raw, targets in mapping.items():
        total = sum(w for _, w in targets)
        if abs(total - 1) > 1e-6:
            raise SystemExit(f"{path}: веса для {raw!r} дают {total:g}, нужно 1")
    return mapping


# --------------------------------------------------------------------------
# Разбор листа
# --------------------------------------------------------------------------

def _read_rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(min_col=1, max_col=SHEET_WIDTH, values_only=True)]


def _find_header(rows, start: int, end: int) -> tuple[int, dict[int, str]] | None:
    """Первая строка блока, где в столбцах проектов есть названия."""
    for i in range(start, end):
        names = {}
        for col in range(FIRST_PROJECT_COL, LAST_PROJECT_COL + 1):
            cell = rows[i][col - 1]
            if isinstance(cell, str) and clean(cell) and clean(cell).casefold() not in IGNORED_HEADERS:
                names[col] = clean(cell)
        if names:
            return i, names
    return None


def _is_total_row(name: str | None, formulas: list) -> bool:
    if name and (key(name).startswith("итого") or key(name) == "без бэка"):
        return True
    return any(isinstance(formulas[c - 1], str) and formulas[c - 1].upper().startswith("=SUM(")
               for c in range(FIRST_PROJECT_COL, LAST_PROJECT_COL + 1))


def _same_shares(a: dict[str, float], b: dict[str, float]) -> bool:
    keys = set(a) | set(b)
    return all(abs(a.get(k, 0) - b.get(k, 0)) < 1e-9 for k in keys)


def _read_block(res: Result, rows_v, rows_f, start: int, end: int,
                month: str, forecast: bool) -> list[SourceRow]:
    header = _find_header(rows_v, start, end)
    if header is None:
        res.add(ERROR, month, start, "в блоке не найдена шапка с названиями проектов")
        return []
    header_idx, projects = header
    res.raw_header[month] = [projects[c] for c in sorted(projects)]
    out: list[SourceRow] = []

    for i in range(header_idx + 1, end):
        row, frow, xl_row = rows_v[i], rows_f[i], i + 1
        name_cell = row[COL_NAME - 1]
        name = clean(name_cell) if isinstance(name_cell, str) and clean(name_cell) else None
        if _is_total_row(name, frow):
            continue
        if name and key(name) == "фио":
            continue

        shares: dict[str, float] = {}
        bad_cell = False
        for col in range(FIRST_PROJECT_COL, LAST_PROJECT_COL + 1):
            cell = f"{get_column_letter(col)}{xl_row}"
            try:
                share = parse_share(row[col - 1], frow[col - 1])
            except ValueError as exc:
                res.add(ERROR, month, xl_row, f"{cell}: {exc}")
                bad_cell = True
                continue
            if not share:
                continue
            if col not in projects:
                res.add(ERROR, month, xl_row, f"{cell}: доля {share:g} в столбце без названия проекта")
                bad_cell = True
                continue
            shares[projects[col]] = share

        if name is None:
            if shares or bad_cell:
                res.add(ERROR, month, xl_row, "есть доли, но не указано ФИО, строка не загружена")
            continue
        if bad_cell:
            res.add(ERROR, month, xl_row, f"{name}: строка не загружена из-за ошибок выше")
            continue

        tag_cell = row[COL_TAG - 1]
        tag = clean(tag_cell) if isinstance(tag_cell, str) else ""
        pos_cell = row[COL_POSITION - 1]
        position = clean(pos_cell) if isinstance(pos_cell, str) else ""
        hire_cell = row[COL_HIRE - 1]
        hire = hire_cell.date() if isinstance(hire_cell, datetime) else None
        if hire_cell is not None and hire is None and clean(hire_cell):
            res.add(WARNING, month, xl_row, f"{name}: в столбце даты найма не дата: {clean(hire_cell)!r}")
        out.append(SourceRow(month, forecast, xl_row, tag, name, position, hire, shares))
    return out


def parse_workbook(path: Path, aliases: dict[str, str] | None = None,
                   project_map: dict[str, list[tuple[str, float]]] | None = None,
                   sheet: str | None = None) -> Result:
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path)
    name = sheet or wb_v.sheetnames[0]
    rows_v, rows_f = _read_rows(wb_v[name]), _read_rows(wb_f[name])

    res = Result()
    titles = []
    for i, row in enumerate(rows_v):
        parsed = parse_title(row[COL_NAME - 1])
        if parsed:
            titles.append((i, *parsed))
    if not titles:
        raise SystemExit("в файле не найдено ни одного заголовка месяца")

    default_names: dict[str, str] = {}      # ключ проекта -> первое написание
    seen_months: set[str] = set()

    for n, (idx, year, month_no, forecast) in enumerate(titles):
        end = titles[n + 1][0] if n + 1 < len(titles) else len(rows_v)
        if year is None:
            res.add(ERROR, f"?-{month_no:02d}", idx + 1, "в заголовке месяца нет года, блок пропущен")
            continue
        month = f"{year}-{month_no:02d}"
        if month in seen_months:
            res.add(ERROR, month, idx + 1, "месяц встречается в файле второй раз, блок пропущен")
            continue
        seen_months.add(month)
        res.months.append(month)
        if forecast:
            res.add(INFO, month, idx + 1, "блок помечен как прогноз")

        source = _read_block(res, rows_v, rows_f, idx + 1, end, month, forecast)
        _load_block(res, source, month, forecast, idx + 1, aliases, project_map, default_names)

    _suggest(res)
    res.months.sort()
    return res


def _canonical_name(name: str, aliases: dict[str, str] | None) -> str:
    return aliases.get(key(name), name) if aliases else name


def _load_block(res: Result, source: list[SourceRow], month: str, forecast: bool,
                title_row: int, aliases, project_map, default_names: dict[str, str]) -> None:
    for r in source:
        r.name = _canonical_name(r.name, aliases)

    # Сначала обычные строки, потом помеченные "дубль", чтобы дубль не стал основной записью.
    ordered = sorted(source, key=lambda r: r.tag.casefold() == DUPLICATE_TAG)
    accepted: dict[str, SourceRow] = {}
    for r in ordered:
        k = key(r.name)
        if k in accepted:
            first = accepted[k]
            if r.tag.casefold() == DUPLICATE_TAG:
                if _same_shares(first.shares, r.shares):
                    res.add(INFO, month, r.row, f"{r.name}: дубль строки {first.row}, значения совпадают, пропущен")
                else:
                    res.add(WARNING, month, r.row,
                            f"{r.name}: дубль строки {first.row}, но доли отличаются, пропущен. Проверьте, какая строка верна")
            else:
                res.add(ERROR, month, r.row,
                        f"{r.name}: ФИО повторяется (первая строка {first.row}) без пометки \"дубль\", строка не загружена")
            continue
        if r.tag.casefold() == DUPLICATE_TAG:
            res.add(WARNING, month, r.row, f"{r.name}: помечено \"дубль\", но основной строки нет, загружено как есть")
        accepted[k] = r

    # Какие названия проектов использованы в этом блоке
    used_raw = {p for r in accepted.values() for p in r.shares}
    for raw in used_raw:
        res.raw_projects.setdefault(raw, set()).add(month)
        if "+" in raw:
            res.add(INFO, month, title_row, f"проект {raw!r} похож на объединенный, при необходимости разбейте его в project_map")
        if project_map is not None and key(raw) not in project_map:
            res.add(ERROR, month, title_row, f"проект {raw!r} не описан в project_map")

    def targets(raw: str) -> list[tuple[str, float]]:
        if project_map is not None and key(raw) in project_map:
            return project_map[key(raw)]
        return [(default_names.setdefault(key(raw), raw), 1.0)]

    order: list[str] = []
    for raw in res.raw_header.get(month, []):
        for project, _ in targets(raw):
            if project not in order:
                order.append(project)
    res.project_positions[month] = order

    for r in sorted(accepted.values(), key=lambda r: r.row):
        total = sum(r.shares.values())
        k = key(r.name)
        res.names.setdefault(k, r.name)
        display = res.names[k]

        if _looks_suspicious(r.name):
            res.add(WARNING, month, r.row, f"{r.name}: не похоже на ФИО человека (заглушка, примечание или группа?)")
        if NOTE_IN_POSITION_RE.search(r.position.casefold()):
            res.add(WARNING, month, r.row, f"{r.name}: в должности похоже на примечание: {r.position!r}")
        if r.hire and (r.hire.year, r.hire.month) > tuple(map(int, month.split("-"))):
            res.add(WARNING, month, r.row, f"{r.name}: дата найма {r.hire} позже месяца аллокации")
        if total == 0:
            res.add(WARNING, month, r.row, f"{r.name}: нет ни одной доли")
        elif abs(total - 1) > TOLERANCE:
            res.add(ERROR, month, r.row, f"{r.name}: сумма долей {total:g}, должна быть 1")

        res.employee_months.append({
            "month": month, "employee": display, "position": r.position,
            "hire_date": r.hire.isoformat() if r.hire else "", "tag": r.tag,
            "total_share": round(total, 6), "is_forecast": int(forecast), "source_row": r.row})

        merged: dict[str, float] = {}
        for raw, share in r.shares.items():
            for project, weight in targets(raw):
                merged[project] = merged.get(project, 0.0) + share * weight
        for project, share in merged.items():
            res.allocations.append({
                "month": month, "employee": display, "project": project,
                "share": round(share, 6), "is_forecast": int(forecast), "source_row": r.row})


def _looks_suspicious(name: str) -> bool:
    words = name.split()
    return len(words) < 2 or name[0].islower() or any(c.isdigit() or c in "()" for c in name)


# --------------------------------------------------------------------------
# Подсказки по справочникам
# --------------------------------------------------------------------------

def _tokens(name: str) -> set[str]:
    return set(key(re.sub(r"\(.*?\)", " ", name)).split())


def suggest_aliases(names: list[str]) -> list[tuple[str, str, str]]:
    """Пары (короткое написание, полное написание, причина) для проверки человеком."""
    found: dict[tuple[str, str], str] = {}
    for a in names:
        for b in names:
            if key(a) == key(b):
                continue
            ta, tb = _tokens(a), _tokens(b)
            if len(ta) >= 2 and ta <= tb and (len(ta) < len(tb) or len(a) < len(b)):
                found[(a, b)] = "то же ФИО, но короче или с примечанием"
    pairs = {frozenset(p) for p in found}
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            if frozenset((a, b)) in pairs or len(_tokens(a)) < 2 or len(_tokens(b)) < 2:
                continue
            if difflib.SequenceMatcher(None, key(a), key(b)).ratio() >= 0.9:
                short, long_ = sorted((a, b), key=len)
                found[(short, long_)] = "очень похожие написания, возможно опечатка"
    return [(a, b, why) for (a, b), why in sorted(found.items())]


def _suggest(res: Result) -> None:
    for alias, canonical, why in suggest_aliases(list(res.names.values())):
        res.add(INFO, "", 0, f"возможно один человек: {alias!r} и {canonical!r} ({why})")


# --------------------------------------------------------------------------
# Вывод результатов
# --------------------------------------------------------------------------

def write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=DELIM)
        writer.writerow(header)
        writer.writerows(rows)


def num(value: float) -> str:
    return f"{value:g}".replace(".", ",")


SCHEMA = """
CREATE TABLE IF NOT EXISTS employees (
    id INTEGER PRIMARY KEY,
    name TEXT NOT NULL UNIQUE,
    external_id TEXT);
CREATE TABLE IF NOT EXISTS projects (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE);
CREATE TABLE IF NOT EXISTS employee_months (
    month TEXT NOT NULL,
    employee_id INTEGER NOT NULL REFERENCES employees(id),
    position TEXT, hire_date TEXT, tag TEXT,
    total_share REAL NOT NULL,
    is_forecast INTEGER NOT NULL,
    source TEXT NOT NULL,
    source_row INTEGER NOT NULL,
    PRIMARY KEY (month, employee_id));
CREATE TABLE IF NOT EXISTS month_projects (
    month TEXT NOT NULL,
    project_id INTEGER NOT NULL REFERENCES projects(id),
    position INTEGER,
    PRIMARY KEY (month, project_id));
CREATE TABLE IF NOT EXISTS allocations (
    month TEXT NOT NULL,
    employee_id INTEGER NOT NULL,
    project_id INTEGER NOT NULL REFERENCES projects(id),
    share REAL NOT NULL,
    hours REAL,
    PRIMARY KEY (month, employee_id, project_id),
    FOREIGN KEY (month, employee_id) REFERENCES employee_months(month, employee_id));
"""


def check_schema(db: sqlite3.Connection, path: Path) -> None:
    """База, созданная прежней версией скрипта, не содержит новых столбцов."""
    columns = {r[1] for r in db.execute("PRAGMA table_info(employees)")}
    if columns and "external_id" not in columns:
        raise SystemExit(f"{path}: база создана прежней версией скриптов. Пересоздайте ее: "
                         "запустите allocation_import.py с исходным Excel еще раз")


def locked_months(db: sqlite3.Connection) -> set[str]:
    """Месяцы, закрытые в веб-интерфейсе (таблицы может и не быть)."""
    try:
        return {r[0] for r in db.execute("SELECT month FROM month_status WHERE locked = 1")}
    except sqlite3.OperationalError:
        return set()


def save_to_db(res: Result, path: Path) -> None:
    """Записывает месяцы из res в базу. Данные тех же месяцев заменяются, остальные не трогаются."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with contextlib.closing(sqlite3.connect(path)) as db:
        check_schema(db, path)
        blocked = locked_months(db) & set(res.months)
        if blocked:
            raise SystemExit(f"месяцы {', '.join(sorted(blocked))} закрыты в веб-интерфейсе, загрузка их не перезапишет. "
                             "Откройте месяц заново в веб-интерфейсе или уберите его из загрузки")
        db.executescript(SCHEMA)
        for month in res.months:
            db.execute("DELETE FROM allocations WHERE month = ?", (month,))
            db.execute("DELETE FROM employee_months WHERE month = ?", (month,))
            db.execute("DELETE FROM month_projects WHERE month = ?", (month,))

        def ref(table: str, name: str, **extra) -> int:
            db.execute(f"INSERT OR IGNORE INTO {table} (name) VALUES (?)", (name,))
            row_id = db.execute(f"SELECT id FROM {table} WHERE name = ?", (name,)).fetchone()[0]
            for column, value in extra.items():
                if value:
                    db.execute(f"UPDATE {table} SET {column} = ? WHERE id = ? AND {column} IS NULL", (value, row_id))
            return row_id

        employees = {}
        for r in res.employee_months:
            employees[r["employee"]] = ref("employees", r["employee"], external_id=r.get("external_id"))
            db.execute(
                "INSERT INTO employee_months (month, employee_id, position, hire_date, tag, total_share,"
                " is_forecast, source, source_row) VALUES (?,?,?,?,?,?,?,?,?)",
                (r["month"], employees[r["employee"]], r["position"], r["hire_date"] or None,
                 r["tag"] or None, r["total_share"], r["is_forecast"], r.get("source", "excel"), r["source_row"]))
        for r in res.allocations:
            db.execute(
                "INSERT INTO allocations (month, employee_id, project_id, share, hours) VALUES (?,?,?,?,?)",
                (r["month"], employees[r["employee"]], ref("projects", r["project"]), r["share"], r.get("hours")))
        for month, names in res.project_positions.items():
            for position, name in enumerate(names):
                db.execute("INSERT OR REPLACE INTO month_projects (month, project_id, position) VALUES (?,?,?)",
                           (month, ref("projects", name), position))
        db.commit()


def has_web_edits(path: Path) -> bool:
    if not path.exists():
        return False
    with contextlib.closing(sqlite3.connect(path)) as db:
        try:
            return db.execute("SELECT COUNT(*) FROM audit_log").fetchone()[0] > 0
        except sqlite3.OperationalError:
            return False


def write_sqlite(res: Result, path: Path) -> None:
    """Пересоздает базу с нуля (для разового импорта из Excel)."""
    if has_web_edits(path):
        raise SystemExit(f"{path}: в базе есть правки, внесенные через веб-интерфейс, пересоздание их удалит. "
                         "Укажите другую папку результатов (--out) или удалите файл базы вручную, если это нужно")
    path.unlink(missing_ok=True)
    save_to_db(res, path)


def write_report(res: Result, path: Path) -> None:
    counts = {lvl: sum(i.level == lvl for i in res.issues) for lvl in (ERROR, WARNING, INFO)}
    lines = ["# Отчет об импорте аллокации", "",
             f"- месяцев: {len(res.months)}",
             f"- записей сотрудник-месяц: {len(res.employee_months)}",
             f"- строк аллокации: {len(res.allocations)}",
             f"- ошибок: {counts[ERROR]}, предупреждений: {counts[WARNING]}, заметок: {counts[INFO]}", ""]

    titles = {ERROR: "Ошибки (исправить в исходном файле или в справочниках)",
              WARNING: "Предупреждения (проверить)", INFO: "Заметки"}
    for level in (ERROR, WARNING, INFO):
        items = sorted((i for i in res.issues if i.level == level), key=lambda i: (i.month, i.row))
        if not items:
            continue
        lines += [f"## {titles[level]}", ""]
        for i in items:
            where = f"{i.month}, строка {i.row}" if i.row else (i.month or "общее")
            lines.append(f"- {where}: {i.message}")
        lines.append("")

    fte: dict[str, dict[str, float]] = {}
    for a in res.allocations:
        fte.setdefault(a["project"], {})
        fte[a["project"]][a["month"]] = fte[a["project"]].get(a["month"], 0) + a["share"]
    if fte:
        lines += ["## Загрузка по проектам, ставок в месяц", "",
                  "| Проект | " + " | ".join(res.months) + " |",
                  "|---|" + "---|" * len(res.months)]
        for project in sorted(fte):
            cells = [f"{fte[project][m]:.2f}" if m in fte[project] else "" for m in res.months]
            lines.append(f"| {project} | " + " | ".join(cells) + " |")
        people = {m: sum(r["month"] == m for r in res.employee_months) for m in res.months}
        lines += ["| Человек в месяце | " + " | ".join(str(people[m]) for m in res.months) + " |", ""]

    path.write_text("\n".join(lines), encoding="utf-8")


def write_outputs(res: Result, out: Path, db: Path | None = None) -> None:
    """Пишет CSV и отчет в out. База: по умолчанию out/allocation.sqlite пересоздается,
    а явно заданная db обновляется по месяцам, не стирая остальные."""
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / "allocations.csv",
               ["month", "employee", "project", "share", "is_forecast", "source_row"],
               [[r["month"], r["employee"], r["project"], num(r["share"]), r["is_forecast"], r["source_row"]]
                for r in res.allocations])
    write_csv(out / "employee_months.csv",
               ["month", "employee", "position", "hire_date", "tag", "total_share", "is_forecast", "source_row"],
               [[r["month"], r["employee"], r["position"], r["hire_date"], r["tag"],
                 num(r["total_share"]), r["is_forecast"], r["source_row"]] for r in res.employee_months])
    write_csv(out / "project_map.suggested.csv", ["raw_name", "project", "weight"],
               [[raw, raw, "1"] for raw in sorted(res.raw_projects, key=lambda p: min(res.raw_projects[p]))])
    write_csv(out / "employee_aliases.suggested.csv", ["alias", "canonical"],
               [[a, b] for a, b, _ in suggest_aliases(list(res.names.values()))])
    if db is None:
        write_sqlite(res, out / "allocation.sqlite")
    else:
        save_to_db(res, db)
    write_report(res, out / "report.md")


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", type=Path, help="исходный файл с аллокацией")
    p.add_argument("--out", type=Path, default=Path("out"), help="папка для результатов (по умолчанию out)")
    p.add_argument("--sheet", help="имя листа (по умолчанию первый)")
    p.add_argument("--project-map", type=Path, default=Path("config/project_map.csv"),
                   help="справочник проектов (используется, если файл существует)")
    p.add_argument("--aliases", type=Path, default=Path("config/employee_aliases.csv"),
                   help="справочник написаний ФИО (используется, если файл существует)")
    p.add_argument("--strict", action="store_true", help="завершиться с кодом 1, если есть ошибки")
    args = p.parse_args(argv)

    aliases = load_aliases(args.aliases) if args.aliases.exists() else None
    project_map = load_project_map(args.project_map) if args.project_map.exists() else None
    res = parse_workbook(args.xlsx, aliases, project_map, args.sheet)
    write_outputs(res, args.out)

    errors = sum(i.level == ERROR for i in res.issues)
    warnings = sum(i.level == WARNING for i in res.issues)
    print(f"Месяцев: {len(res.months)}, записей сотрудник-месяц: {len(res.employee_months)}, "
          f"строк аллокации: {len(res.allocations)}")
    print(f"Ошибок: {errors}, предупреждений: {warnings}. Подробности: {args.out / 'report.md'}")
    if aliases is None:
        print(f"Справочник ФИО не задан ({args.aliases}), написания не объединялись.")
    if project_map is None:
        print(f"Справочник проектов не задан ({args.project_map}), названия взяты как в файле.")
    return 1 if args.strict and errors else 0


if __name__ == "__main__":
    sys.exit(main())
