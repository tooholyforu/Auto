#!/usr/bin/env python3
"""Расчет расходов на сотрудников по проектам за месяц (блок "распределение по проектам" из ТЗ).

Расходы сотрудника умножаются на его доли по проектам (аллокация), потом суммируются по проектам:
- по начислению: оклады начислено, страховые начислено;
- по оплате: оклады оплачено, НДФЛ, страховые оплаченные;
- сотрудники: количество человек и ставок (доля x количество ставок) только по тем, кто есть в 1С.
  Внешние исполнители (нет в 1С) считаются отдельной строкой.

Расходы людей без аллокации, а также нераспределенный остаток при сумме долей меньше 1 показываются
отдельной колонкой "Не распределено", чтобы итог сходился с общими расходами.

Данные: аллокация из базы allocation.sqlite, расходы из базы costs.sqlite (onec_import.py).
По умолчанию в Excel попадает только сводка по проектам. Строки по сотрудникам (с зарплатами) только с --detail.

    python cost_report.py --month 2026-09
"""
from __future__ import annotations

import argparse
import contextlib
import sqlite3
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from allocation_import import ERROR, INFO, WARNING, Issue, key, load_aliases
from reconcile_jira import read_allocations

MONEY_FIELDS = ["accrued", "insurance_accrued", "paid", "ndfl", "insurance_paid"]
UNDISTRIBUTED = "Не распределено"


@dataclass
class Summary:
    month: str
    projects: list[str]
    totals: dict[str, dict[str, float]] = field(default_factory=dict)    # колонка -> показатель -> сумма
    issues: list[Issue] = field(default_factory=list)
    people: list[dict] = field(default_factory=list)                     # строки по сотрудникам для --detail


def read_costs(db_path: Path, month: str) -> list[dict]:
    with contextlib.closing(sqlite3.connect(db_path)) as db:
        db.row_factory = sqlite3.Row
        return [dict(r) for r in db.execute("SELECT * FROM employee_costs WHERE month = ?", (month,))]


def compute(month: str, allocations: list[dict], costs: list[dict], aliases: dict[str, str] | None = None,
            headcount_exclude: set[str] | None = None) -> Summary:
    aliases = aliases or {}
    exclude = {key(n) for n in (headcount_exclude or set())}
    cost_by_key = {key(c["employee"]): c for c in costs}

    shares: dict[str, dict[str, float]] = {}
    display: dict[str, str] = {}
    for a in allocations:
        name = aliases.get(key(a["employee"]), a["employee"])
        shares.setdefault(key(name), {})
        shares[key(name)][a["project"]] = shares[key(name)].get(a["project"], 0.0) + a["share"]
        display.setdefault(key(name), name)

    projects = sorted({p for s in shares.values() for p in s})
    out = Summary(month, projects)
    metrics = MONEY_FIELDS + ["headcount", "fte", "outside"]
    out.totals = {p: dict.fromkeys(metrics, 0.0) for p in projects}
    grand = dict.fromkeys(MONEY_FIELDS, 0.0)
    for c in costs:
        for f in MONEY_FIELDS:
            grand[f] += c[f]

    distributed = dict.fromkeys(MONEY_FIELDS, 0.0)
    no_cost, over = [], []
    for k, per_project in shares.items():
        c = cost_by_key.get(k)
        total_share = sum(per_project.values())
        if total_share > 1 + 0.005:
            over.append(f"{display[k]} ({total_share:g})")
        if c is None:
            no_cost.append(display[k])
        rate = (c or {}).get("rate") or 0.0
        for project, share in per_project.items():
            slot = out.totals[project]
            if c is not None:
                for f in MONEY_FIELDS:
                    slot[f] += c[f] * share
                    distributed[f] += c[f] * share
            if c is None:
                slot["outside"] += share           # нет данных 1С: внешний исполнитель
            else:
                if k not in exclude:
                    slot["headcount"] += share
                slot["fte"] += share * rate        # исключение из количества на ставки не действует (как в ТЗ)
        if c is not None:
            out.people.append({"name": c["employee"], "position": c.get("position") or "", "department": c.get("department") or "",
                               "rate": c.get("rate"), "base_salary": c.get("base_salary"),
                               **{f: c[f] for f in MONEY_FIELDS}, "shares": per_project})

    rest = {f: grand[f] - distributed[f] for f in MONEY_FIELDS}
    out.totals[UNDISTRIBUTED] = {**rest, "headcount": 0.0, "fte": 0.0, "outside": 0.0}
    unallocated = sorted(c["employee"] for c in costs if key(c["employee"]) not in shares)
    for c in costs:
        if key(c["employee"]) not in shares:
            out.people.append({"name": c["employee"], "position": c.get("position") or "", "department": c.get("department") or "",
                               "rate": c.get("rate"), "base_salary": c.get("base_salary"),
                               **{f: c[f] for f in MONEY_FIELDS}, "shares": {}})

    if unallocated:
        out.issues.append(Issue(ERROR, month, 0, f"расходы есть, а аллокации нет ({len(unallocated)}): {', '.join(unallocated)}. "
                                                 "Их расходы попали в колонку \"Не распределено\". Добавьте их в аллокацию "
                                                 "(например, уволенные в течение месяца) или проверьте написание ФИО"))
    if no_cost:
        out.issues.append(Issue(WARNING, month, 0, f"в аллокации есть, а расходов из 1С нет ({len(no_cost)}): {', '.join(sorted(no_cost))}. "
                                                   "Скорее всего, внешние исполнители. Их расходы не учтены"))
    if over:
        out.issues.append(Issue(ERROR, month, 0, f"сумма долей больше 1: {', '.join(over)}, расходы посчитаны дважды"))
    short = [display[k] for k, s in shares.items() if k in cost_by_key and sum(s.values()) < 1 - 0.005]
    if short:
        out.issues.append(Issue(WARNING, month, 0, f"сумма долей меньше 1: {', '.join(sorted(short))}. Остаток расходов в колонке \"Не распределено\""))
    return out


ROWS = [
    ("по начислению", None), ("Оклады начислено", "accrued"), ("Страховые начислено", "insurance_accrued"),
    ("Итого начислено", ("accrued", "insurance_accrued")),
    ("по оплате", None), ("Оклады оплачено", "paid"), ("НДФЛ оплачен", "ndfl"), ("Страховые оплаченные", "insurance_paid"),
    ("Итого оплачено (оклады, НДФЛ, страховые)", ("paid", "ndfl", "insurance_paid")),
    ("сотрудники", None), ("Количество сотрудников", "headcount"), ("Ставок (доля x ставки)", "fte"),
    ("Внешние исполнители без данных 1С (человек)", "outside"),
]


def write_excel(s: Summary, path: Path, detail: bool = False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "По проектам"
    columns = s.projects + [UNDISTRIBUTED]
    ws.append([f"Расходы на сотрудников по проектам, {s.month}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([""] + columns + ["Итого"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for label, spec in ROWS:
        row = ws.max_row + 1
        ws.cell(row, 1, label)
        if spec is None:
            ws.cell(row, 1).font = Font(bold=True, italic=True)
            continue
        names = (spec,) if isinstance(spec, str) else spec
        for i, col in enumerate(columns, 2):
            ws.cell(row, i, round(sum(s.totals[col][n] for n in names), 2))
        total_col = len(columns) + 2
        ws.cell(row, total_col, f"=SUM({get_column_letter(2)}{row}:{get_column_letter(total_col - 1)}{row})")
        fmt = "#,##0.00" if names[0] not in ("headcount", "fte", "outside") else "0.00"
        for i in range(2, total_col + 1):
            ws.cell(row, i).number_format = fmt
        if len(names) > 1 or label.startswith("Итого"):
            for i in range(1, total_col + 1):
                ws.cell(row, i).font = Font(bold=True)
    ws.column_dimensions["A"].width = 44
    for i in range(2, len(columns) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "B4"

    if detail:
        wd = wb.create_sheet("Сотрудники")
        wd.append(["ФИО", "Должность", "Подразделение", "Ставок", "Оклад на ставку", "Начислено", "Страховые начислено",
                   "Оплачено", "НДФЛ", "Страховые оплачено"] + s.projects)
        for cell in wd[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for p in sorted(s.people, key=lambda p: p["name"]):
            wd.append([p["name"], p["position"], p["department"], p["rate"], p["base_salary"], p["accrued"],
                       p["insurance_accrued"], p["paid"], p["ndfl"], p["insurance_paid"]]
                      + [p["shares"].get(project) for project in s.projects])
        wd.column_dimensions["A"].width = 38
        wd.column_dimensions["B"].width = 36
        wd.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    with contextlib.suppress(OSError):
        path.chmod(0o600)


def write_notes(s: Summary, path: Path) -> None:
    lines = [f"# Расчет расходов по проектам за {s.month}", ""]
    titles = {ERROR: "Ошибки", WARNING: "Предупреждения", INFO: "Заметки"}
    for level in (ERROR, WARNING, INFO):
        items = [i for i in s.issues if i.level == level]
        if items:
            lines += [f"## {titles[level]}", ""] + [f"- {i.message}" for i in items] + [""]
    if len(lines) == 2:
        lines.append("Замечаний нет.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--month", required=True, help="месяц ГГГГ-ММ")
    p.add_argument("--alloc-db", type=Path, default=Path("out/allocation.sqlite"), help="база с аллокацией")
    p.add_argument("--costs-db", type=Path, default=Path("out/costs.sqlite"), help="база расходов из 1С")
    p.add_argument("--aliases", type=Path, default=Path("config/employee_aliases.csv"),
                   help="написания ФИО (используется, если файл существует)")
    p.add_argument("--headcount-exclude", nargs="*", default=[], metavar="ФИО",
                   help="сотрудники, которых не учитывать в количестве человек (расходы и ставки сохраняются)")
    p.add_argument("--headcount-exclude-file", type=Path, default=Path("config/headcount_exclude.txt"),
                   help="то же списком, по одному ФИО в строке (используется, если файл существует)")
    p.add_argument("--detail", action="store_true", help="добавить в Excel лист со строками по сотрудникам (зарплаты!)")
    p.add_argument("--out", type=Path, default=Path("out/costs"), help="папка для результатов")
    args = p.parse_args(argv)

    for path, what in ((args.alloc_db, "база с аллокацией"), (args.costs_db, "база расходов")):
        if not path.exists():
            raise SystemExit(f"не найден файл: {what} ({path})")
    allocations = read_allocations(args.alloc_db, args.month)
    costs = read_costs(args.costs_db, args.month)
    if not allocations:
        raise SystemExit(f"в {args.alloc_db} нет аллокации за {args.month}")
    if not costs:
        raise SystemExit(f"в {args.costs_db} нет расходов за {args.month}")

    aliases = load_aliases(args.aliases) if args.aliases.exists() else None
    exclude = set(args.headcount_exclude)
    if args.headcount_exclude_file.exists():
        exclude |= {line.strip() for line in args.headcount_exclude_file.read_text(encoding="utf-8").splitlines() if line.strip()}
    summary = compute(args.month, allocations, costs, aliases, exclude)
    write_excel(summary, args.out / f"projects-{args.month}.xlsx", args.detail)
    write_notes(summary, args.out / f"projects-{args.month}-notes.md")
    errors = sum(i.level == ERROR for i in summary.issues)
    print(f"{args.month}: проектов {len(summary.projects)}, ошибок {errors}, "
          f"предупреждений {sum(i.level == WARNING for i in summary.issues)}. "
          f"Excel: {args.out / f'projects-{args.month}.xlsx'}, замечания: {args.out / f'projects-{args.month}-notes.md'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
