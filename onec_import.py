#!/usr/bin/env python3
"""Загрузка выгрузок 1С для расчета расходов на сотрудников по месяцам.

Читает шесть отчетов (xlsx или xls) и складывает по каждому сотруднику суммы в базу расходов:

| Поле базы           | Отчет 1С                                              | Столбец                      |
|---------------------|-------------------------------------------------------|------------------------------|
| должность, ставка,  | Штатные сотрудники                                    | Должность, Количество ставок,|
| оклад на ставку     |                                                       | Оклад, коэффициенты          |
| accrued             | Анализ зарплаты по сотрудникам (помесячно)            | Всего начислено              |
| insurance_accrued   | Налоги и взносы (кратко) за месяц расчета             | Единый тариф + Несч. случаи  |
| paid                | Карточка счета 51 со статьей "Оклады"                 | сумма по кредиту, по ФИО     |
| ndfl                | Анализ НДФЛ по месяцам                                | Исчислено налога             |
| insurance_paid      | Налоги и взносы (кратко) за предыдущий месяц          | Единый тариф + Несч. случаи  |

Столбцы ищутся по названиям в шапке, а не по буквам. В базе лежат зарплаты, поэтому база
расходов хранится отдельно от базы аллокации и должна быть доступна только финансам.

Пример: все шесть файлов лежат в одной папке, имена не важны, отчеты узнаются по содержимому:
    python onec_import.py --month 2026-09 --dir data/1c
Или файлы по одному: --staff, --payroll, --taxes, --taxes-prev, --ndfl, --card51.
"""
from __future__ import annotations

import argparse
import contextlib
import re
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from allocation_import import ERROR, INFO, WARNING, Issue, clean, key, parse_title

NON_PERSON_RE = re.compile(r"(^|\s)(ООО|АО|ПАО|ЗАО|ИП)(\s|$)|^итого|^подразделение|^сотрудник$|^организация|^обороты",
                           re.IGNORECASE)
DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")
TOLERANCE = 0.5          # рубли: допуск при сверке с итоговой строкой отчета


# --------------------------------------------------------------------------
# Чтение листов
# --------------------------------------------------------------------------

def read_sheet(path: Path) -> list[list]:
    """Первый лист файла как список строк. Пустые строки и ячейки превращаются в None."""
    if path.suffix.lower() == ".xls":
        import xlrd
        sheet = xlrd.open_workbook(path).sheet_by_index(0)
        rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=0)
    return [[None if v is None or (isinstance(v, str) and not v.strip()) else v for v in r] + [None] * (width - len(r))
            for r in rows]


def number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(clean(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def find_header(rows: list[list], labels: dict[str, str], where: str, limit: int = 40) -> tuple[int, dict[str, int]]:
    """Находит строку шапки, где есть все столбцы, и возвращает индексы столбцов по ключам."""
    wanted = {k: v.casefold() for k, v in labels.items()}
    for i, row in enumerate(rows[:limit]):
        cells = {c: key(v) for c, v in enumerate(row) if isinstance(v, str)}
        found = {}
        for name, label in wanted.items():
            col = next((c for c, text in cells.items() if text.startswith(label)), None)
            if col is not None:
                found[name] = col
        if len(found) == len(wanted):
            return i, found
    missing = ", ".join(labels.values())
    raise SystemExit(f"{where}: в шапке не найдены столбцы: {missing}")


def is_person(name, non_people: set[str]) -> bool:
    return (isinstance(name, str) and len(clean(name).split()) >= 2 and not NON_PERSON_RE.search(clean(name))
            and key(name) not in non_people)


def month_of(value) -> str | None:
    parsed = parse_title(value) if isinstance(value, (str, datetime)) else None
    return f"{parsed[0]}-{parsed[1]:02d}" if parsed and parsed[0] else None


def prev_month(month: str) -> str:
    year, number_ = map(int, month.split("-"))
    return f"{year - (number_ == 1)}-{(number_ - 2) % 12 + 1:02d}"


# --------------------------------------------------------------------------
# Разбор отчетов
# --------------------------------------------------------------------------

@dataclass
class StaffRow:
    row: int
    name: str
    tab_no: str
    position: str
    department: str
    rate: float
    base_salary: float | None
    hire: date | None


@dataclass
class Totals:
    values: dict[str, float] = field(default_factory=dict)       # ключ ФИО -> сумма
    names: dict[str, str] = field(default_factory=dict)          # ключ ФИО -> написание в отчете
    rows: dict[str, int] = field(default_factory=dict)
    control: float | None = None                                 # итоговая строка отчета, если нашли
    period: str | None = None

    def add(self, name: str, amount: float, row: int) -> None:
        k = key(name)
        self.values[k] = self.values.get(k, 0.0) + amount
        self.names.setdefault(k, clean(name))
        self.rows.setdefault(k, row)


def parse_rate(value) -> float | None:
    """Количество ставок: 1, 0,4, 2/5 или текст."""
    n = number(value)
    if n is not None:
        return n
    text = clean(value or "")
    m = re.fullmatch(r"(\d+(?:[.,]\d+)?)\s*/\s*(\d+(?:[.,]\d+)?)", text)
    if m and float(m[2].replace(",", ".")):
        return float(m[1].replace(",", ".")) / float(m[2].replace(",", "."))
    return None


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(clean(value or ""), fmt).date()
        except ValueError:
            pass
    return None


def read_staff(path: Path, issues: list[Issue], month: str) -> list[StaffRow]:
    """Штатные сотрудники. Подразделения идут отдельными строками без табельного номера."""
    rows = read_sheet(path)
    h, col = find_header(rows, {"name": "Сотрудник", "tab": "Табельный номер", "position": "Должность",
                                "salary": "Оклад", "rate": "Количество ставок"}, path.name)
    extra = {}
    for name, label in (("hire", "Дата приема"), ("district", "Территория.Районный коэффициент"),
                        ("north", "Процент северной надбавки")):
        for c, v in enumerate(rows[h]):
            if isinstance(v, str) and key(v).startswith(label.casefold()):
                extra[name] = c
    result, department = [], ""
    for i in range(h + 1, len(rows)):
        row = rows[i]
        label = row[0]
        if isinstance(label, str) and clean(label) and row[col["name"]] is None:
            department = clean(label)
            continue
        name, tab = row[col["name"]], row[col["tab"]]
        if not isinstance(name, str) or tab is None:
            continue
        rate = parse_rate(row[col["rate"]])
        if rate is None:
            issues.append(Issue(WARNING, month, i + 1, f"{clean(name)}: не удалось прочитать количество ставок "
                                                       f"{row[col['rate']]!r}, принято 1"))
            rate = 1.0
        base = number(row[col["salary"]])
        if base is not None:
            ak = number(row[extra["district"]]) if "district" in extra else None
            al = number(row[extra["north"]]) if "north" in extra else None
            if ak is not None or al is not None:
                # правило ТЗ: оклад * районный коэффициент + оклад * северная надбавка / 100
                base = base * (ak if ak is not None else 1.0) + base * (al or 0.0) / 100
        result.append(StaffRow(i + 1, clean(name), clean(tab), clean(row[col["position"]] or ""), department, rate, base,
                               parse_date(row[extra["hire"]]) if "hire" in extra else None))
    if not result:
        raise SystemExit(f"{path.name}: не найдено ни одного сотрудника с табельным номером")
    return result


def _check_period(found: str | None, expected: str, what: str, issues: list[Issue], month: str) -> None:
    if found is None:
        issues.append(Issue(WARNING, month, 0, f"{what}: не удалось определить месяц отчета, проверьте вручную"))
    elif found != expected:
        issues.append(Issue(ERROR, month, 0, f"{what}: в отчете месяц {found}, а нужен {expected}"))


def read_payroll(path: Path, non_people: set[str], month: str, issues: list[Issue]) -> Totals:
    """Анализ зарплаты по сотрудникам. Строки подразделений содержат суммы по отделу, их пропускаем."""
    rows = read_sheet(path)
    h, col = find_header(rows, {"total": "Всего начислено"}, path.name)
    out = Totals()
    for i, row in enumerate(rows):
        if i <= h:
            continue
        name = row[0]
        amount = number(row[col["total"]])
        if isinstance(name, str) and key(name).startswith("итого"):
            out.control = amount
        elif is_person(name, non_people) and amount is not None:
            if key(name) in out.values:
                issues.append(Issue(INFO, month, i + 1, f"{clean(name)}: в отчете несколько строк (перевод между отделами?), суммы сложены"))
            out.add(name, amount, i + 1)
    for row in rows[:h + 1]:
        for c, v in enumerate(row):
            if isinstance(v, str) and key(v) == "месяц":
                out.period = next((month_of(x) for x in row[c + 1:] if x is not None), None)
    _check_period(out.period, month, f"{path.name} (зарплата)", issues, month)
    return out


def taxes_period(rows: list[list]) -> str | None:
    """Месяц отчета "Налоги и взносы (кратко)": значение под подписью "Период формирования отчета"."""
    for i, row in enumerate(rows[:12]):
        if isinstance(row[0], str) and key(row[0]).startswith("период формирования"):
            return next((month_of(v) for r in rows[i:i + 3] for v in r if v is not None and month_of(v)), None)
    return None


def read_taxes(path: Path, non_people: set[str], expected: str, month: str, issues: list[Issue], what: str) -> Totals:
    """Налоги и взносы (кратко): единый тариф плюс несчастные случаи."""
    rows = read_sheet(path)
    h, col = find_header(rows, {"tariff": "Единый тариф", "injury": "Несч. случаи"}, path.name)
    out = Totals()
    for i, row in enumerate(rows):
        if i <= h:
            continue
        name = row[0]
        if not isinstance(name, str):
            continue
        amount = (number(row[col["tariff"]]) or 0.0) + (number(row[col["injury"]]) or 0.0)
        if key(name).startswith("итого"):
            out.control = amount
        elif is_person(name, non_people) and not re.search(r"(ООО|АО|ПАО|ИП)\b", name):
            out.add(name, amount, i + 1)
    out.period = taxes_period(rows)
    _check_period(out.period, expected, f"{path.name} ({what})", issues, month)
    return out


def read_ndfl(path: Path, non_people: set[str], month: str, issues: list[Issue],
              months: list[str] | None) -> tuple[Totals, dict[str, int]]:
    """Анализ НДФЛ по месяцам: исчисленный налог. В файле несколько месячных секций, берем нужные."""
    rows = read_sheet(path)
    h, col = find_header(rows, {"tax": "Исчислено налога"}, path.name)
    wanted = months or [month]
    out, sizes, section = Totals(), {}, None
    for i, row in enumerate(rows):
        if i <= h:
            continue
        name = row[0]
        if not isinstance(name, str):
            continue
        marker = month_of(name)
        if marker:
            section = marker
            sizes.setdefault(section, 0)
            continue
        amount = number(row[col["tax"]])
        if key(name).startswith("итого"):
            out.control = amount
        elif section and is_person(name, non_people) and amount is not None:
            sizes[section] += 1
            if section in wanted:
                out.add(name, amount, i + 1)
    missing = [m for m in wanted if m not in sizes]
    if missing:
        issues.append(Issue(ERROR, month, 0, f"{path.name}: в отчете нет секции за {', '.join(missing)}, есть: {', '.join(sizes) or 'ни одной'}"))
    elif len(sizes) > len(wanted):
        others = ", ".join(f"{m} ({n} чел.)" for m, n in sizes.items() if m not in wanted)
        issues.append(Issue(INFO, month, 0, f"{path.name}: учтены секции {', '.join(wanted)}, остальные пропущены: {others}"))
    if out.control is not None and len(sizes) > len(wanted):
        out.control = None            # итог файла включает пропущенные секции
    return out, sizes


def read_card51(path: Path, month: str, issues: list[Issue]) -> Totals:
    """Карточка счета 51 по статье "Оклады": суммы выплат по кредиту, получатель в аналитике дебета.

    Столбец с суммой определяется по расположению: первая числовая ячейка правее счета 51.
    """
    rows = read_sheet(path)
    h, col = find_header(rows, {"date": "Период", "doc": "Документ", "debit_analytics": "Аналитика Дт"}, path.name)
    out = Totals()
    amount_cols: set[int] = set()
    for i, row in enumerate(rows):
        if i <= h:
            continue
        stamp = row[0]
        text = clean(stamp) if stamp is not None else ""
        if key(text).startswith("обороты за период"):
            tail = [number(v) for v in row[3:] if number(v) is not None]
            out.control = tail[1] if len(tail) > 1 else (tail[0] if tail else None)
            continue
        if not DATE_RE.match(text):
            continue
        analytics = row[col["debit_analytics"]]
        credit_account = next((c for c, v in enumerate(row[4:], 4) if clean(v or "") == "51"), None)
        if not isinstance(analytics, str) or credit_account is None:
            continue
        amount_col = next((c for c in range(credit_account + 1, len(row)) if number(row[c]) is not None), None)
        if amount_col is None:
            continue
        amount_cols.add(amount_col)
        out.add(analytics.split("\n")[0], number(row[amount_col]) or 0.0, i + 1)
    title = next((clean(v) for r in rows[:6] for v in r if isinstance(v, str) and "Карточка счета" in v), "")
    m = re.search(r"за\s+([А-Яа-я]+)\s+(\d{4})", title)
    out.period = month_of(f"{m[1]} {m[2]}") if m else None
    _check_period(out.period, month, f"{path.name} (карточка счета 51)", issues, month)
    if amount_cols:
        out_cols = ", ".join(openpyxl.utils.get_column_letter(c + 1) for c in sorted(amount_cols))
        issues.append(Issue(INFO, month, 0, f"{path.name}: суммы выплат взяты из столбца {out_cols}"))
    # контрольная строка карточки содержит итоговые обороты по кредиту
    return out


# --------------------------------------------------------------------------
# Автоопределение отчетов в папке
# --------------------------------------------------------------------------

REPORT_NAMES = {"staff": "Штатные сотрудники", "payroll": "Анализ зарплаты по сотрудникам",
                "taxes": "Налоги и взносы (кратко) за месяц расчета",
                "taxes_prev": "Налоги и взносы (кратко) за предыдущий месяц",
                "ndfl": "Анализ НДФЛ по месяцам", "card51": "Карточка счета 51"}
REPORT_MARKERS = [("staff", "штатные сотрудники"), ("payroll", "анализ зарплаты по сотрудникам"),
                  ("taxes", "налоги и взносы"), ("ndfl", "анализ ндфл"), ("card51", "карточка счета 51")]


def detect_reports(folder: Path, month: str) -> tuple[dict[str, Path], list[str]]:
    """Узнает выгрузки 1С по заголовку внутри файла, имя файла не важно.

    Два отчета "Налоги и взносы" различаются месяцем: за месяц расчета и за предыдущий.
    Возвращает найденные файлы по ролям и список замечаний (лишние и неопознанные файлы).
    """
    if not folder.is_dir():
        raise SystemExit(f"{folder}: не папка или ее нет")
    found: dict[str, list[Path]] = {}
    notes: list[str] = []
    for path in sorted(folder.iterdir()):
        if path.suffix.lower() not in (".xlsx", ".xls") or path.name.startswith(("~$", ".")):
            continue
        try:
            rows = read_sheet(path)
        except Exception:        # файл поврежден или это не Excel: не мешаем остальным
            notes.append(f"{path.name}: файл не удалось прочитать, пропущен")
            continue
        header = [key(v) for r in rows[:10] for v in r if isinstance(v, str)]
        role = next((r for r, marker in REPORT_MARKERS if any(marker in t for t in header)), None)
        if role == "taxes":
            period = taxes_period(rows)
            if period == month:
                role = "taxes"
            elif period == prev_month(month):
                role = "taxes_prev"
            else:
                notes.append(f"{path.name}: отчет «Налоги и взносы» за {period or 'неизвестный месяц'}, "
                             f"а нужны {prev_month(month)} и {month}, пропущен")
                continue
        if role is None:
            notes.append(f"{path.name}: не похож на выгрузку 1С для расчета, пропущен")
            continue
        found.setdefault(role, []).append(path)
    clashes = [f"«{REPORT_NAMES[r]}»: {', '.join(p.name for p in paths)}" for r, paths in found.items() if len(paths) > 1]
    if clashes:
        raise SystemExit("в папке несколько файлов подходят под один отчет, оставьте по одному:\n  " + "\n  ".join(clashes))
    return {r: paths[0] for r, paths in found.items()}, notes


# --------------------------------------------------------------------------
# Сборка по месяцу
# --------------------------------------------------------------------------

@dataclass
class CostRecord:
    name: str
    tab_no: str
    position: str
    department: str
    in_staff: bool
    rate: float | None
    base_salary: float | None
    accrued: float
    insurance_accrued: float
    paid: float
    ndfl: float
    insurance_paid: float


@dataclass
class CostResult:
    month: str
    records: list[CostRecord] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    sources: dict[str, str] = field(default_factory=dict)


def collect_costs(month: str, staff: Path, payroll: Path, taxes: Path, taxes_prev: Path, ndfl: Path, card51: Path,
                  ndfl_months: list[str] | None = None) -> CostResult:
    res = CostResult(month)
    issues = res.issues
    staff_rows = read_staff(staff, issues, month)
    non_people = {key(s.department) for s in staff_rows if s.department}
    for row in read_sheet(staff)[:8]:
        non_people |= {key(v) for v in row if isinstance(v, str) and NON_PERSON_RE.search(v)}

    payroll_t = read_payroll(payroll, non_people, month, issues)
    taxes_t = read_taxes(taxes, non_people, month, month, issues, "взносы за месяц расчета")
    prev_t = read_taxes(taxes_prev, non_people, prev_month(month), month, issues, "взносы за предыдущий месяц")
    ndfl_t, _ = read_ndfl(ndfl, non_people, month, issues, ndfl_months)
    paid_t = read_card51(card51, month, issues)

    sources = {"accrued": payroll_t, "insurance_accrued": taxes_t, "insurance_paid": prev_t, "ndfl": ndfl_t, "paid": paid_t}
    labels = {"accrued": "зарплата", "insurance_accrued": "взносы за месяц", "insurance_paid": "взносы за прошлый месяц",
              "ndfl": "НДФЛ", "paid": "карточка 51"}

    # Сверка с итоговой строкой отчета ловит ошибки разбора
    for field_name, t in sources.items():
        if t.control is not None and abs(sum(t.values.values()) - t.control) > TOLERANCE:
            issues.append(Issue(WARNING, month, 0, f"{labels[field_name]}: сумма по сотрудникам {sum(t.values.values()):,.2f} "
                                                   f"не равна итогу отчета {t.control:,.2f}"))

    by_key = {key(s.name): s for s in staff_rows}
    first_day = date(int(month[:4]), int(month[5:]), 1)
    everyone = dict.fromkeys(list(by_key) + [k for t in sources.values() for k in t.values])
    for k in everyone:
        s = by_key.get(k)
        display = s.name if s else next((t.names[k] for t in sources.values() if k in t.names), k)
        rec = CostRecord(display, s.tab_no if s else "", s.position if s else "", s.department if s else "", s is not None,
                         s.rate if s else None, s.base_salary if s else None,
                         *(sources[f].values.get(k, 0.0) for f in ("accrued", "insurance_accrued", "paid", "ndfl", "insurance_paid")))
        res.records.append(rec)
        if s is None:
            where = [labels[f] for f, t in sources.items() if k in t.values]
            issues.append(Issue(INFO, month, 0, f"{display}: нет в штатном списке, но есть в отчетах ({', '.join(where)}). "
                                                "Вероятно, уволен или работает по договору, включите в аллокацию"))
            continue
        for field_name in ("accrued", "insurance_accrued", "paid", "ndfl"):
            if k not in sources[field_name].values:
                issues.append(Issue(WARNING, month, s.row, f"{display}: нет в отчете ({labels[field_name]}), принято 0"))
        if k not in prev_t.values and not (s.hire and s.hire >= first_day):
            issues.append(Issue(WARNING, month, s.row, f"{display}: нет в отчете ({labels['insurance_paid']}), принято 0"))
    return res


# --------------------------------------------------------------------------
# База расходов
# --------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS employee_costs (
    month TEXT NOT NULL,
    employee TEXT NOT NULL,
    tab_no TEXT, position TEXT, department TEXT,
    in_staff INTEGER NOT NULL,
    rate REAL, base_salary REAL,
    accrued REAL NOT NULL,
    insurance_accrued REAL NOT NULL,
    paid REAL NOT NULL,
    ndfl REAL NOT NULL,
    insurance_paid REAL NOT NULL,
    PRIMARY KEY (month, employee));
"""


def save_costs(res: CostResult, path: Path) -> None:
    """Записывает месяц в базу расходов. Данные этого месяца заменяются, остальные остаются."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with contextlib.closing(sqlite3.connect(path)) as db:
        db.executescript(SCHEMA)
        db.execute("DELETE FROM employee_costs WHERE month = ?", (res.month,))
        db.executemany(
            "INSERT INTO employee_costs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(res.month, r.name, r.tab_no or None, r.position or None, r.department or None, int(r.in_staff), r.rate,
              r.base_salary, round(r.accrued, 2), round(r.insurance_accrued, 2), round(r.paid, 2), round(r.ndfl, 2),
              round(r.insurance_paid, 2)) for r in res.records])
        db.commit()
    with contextlib.suppress(OSError):
        path.chmod(0o600)


def write_report(res: CostResult, path: Path) -> None:
    counts = {lvl: sum(i.level == lvl for i in res.issues) for lvl in (ERROR, WARNING, INFO)}
    lines = [f"# Загрузка расходов из 1С за {res.month}", "",
             f"- сотрудников: {len(res.records)}, из них в штатном списке: {sum(r.in_staff for r in res.records)}",
             f"- ошибок: {counts[ERROR]}, предупреждений: {counts[WARNING]}, заметок: {counts[INFO]}", ""]
    titles = {ERROR: "Ошибки", WARNING: "Предупреждения", INFO: "Заметки"}
    for level in (ERROR, WARNING, INFO):
        items = sorted((i for i in res.issues if i.level == level), key=lambda i: (i.row, i.message))
        if items:
            lines += [f"## {titles[level]}", ""] + [f"- {'строка ' + str(i.row) + ': ' if i.row else ''}{i.message}" for i in items] + [""]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--month", required=True, help="месяц расчета ГГГГ-ММ")
    p.add_argument("--dir", type=Path, help="папка с выгрузками 1С: отчеты определяются по содержимому, имена файлов не важны")
    for flag, text in (("staff", "Штатные сотрудники"), ("payroll", "Анализ зарплаты по сотрудникам (помесячно)"),
                       ("taxes", "Налоги и взносы (кратко) за месяц расчета"),
                       ("taxes-prev", "Налоги и взносы (кратко) за предыдущий месяц"),
                       ("ndfl", "Анализ НДФЛ по месяцам"), ("card51", "Карточка счета 51, статья Оклады")):
        p.add_argument(f"--{flag}", type=Path, help=text + " (если не задана --dir)")
    p.add_argument("--ndfl-months", nargs="+", help="секции отчета НДФЛ, которые учитывать (по умолчанию месяц расчета)")
    p.add_argument("--db", type=Path, default=Path("out/costs.sqlite"), help="база расходов (доступ только финансам)")
    p.add_argument("--out", type=Path, default=Path("out/costs"), help="папка для отчета о загрузке")
    p.add_argument("--strict", action="store_true", help="завершиться с кодом 1, если есть ошибки")
    args = p.parse_args(argv)

    paths = {"staff": args.staff, "payroll": args.payroll, "taxes": args.taxes, "taxes_prev": args.taxes_prev,
             "ndfl": args.ndfl, "card51": args.card51}
    if args.dir:
        detected, notes = detect_reports(args.dir, args.month)
        for role, path in detected.items():
            paths[role] = paths[role] or path          # явно заданный файл важнее найденного
        print(f"Отчеты в {args.dir}:")
        for role, name in REPORT_NAMES.items():
            print(f"  {name}: {paths[role].name if paths[role] else 'НЕ НАЙДЕН'}")
        for note in notes:
            print(f"  ({note})")
    missing = [REPORT_NAMES[r] for r, path in paths.items() if path is None]
    if missing:
        raise SystemExit("не хватает отчетов: " + "; ".join(missing) + ". Укажите --dir с папкой или отдельные файлы.")
    res = collect_costs(args.month, paths["staff"], paths["payroll"], paths["taxes"], paths["taxes_prev"], paths["ndfl"],
                        paths["card51"], args.ndfl_months)
    save_costs(res, args.db)
    write_report(res, args.out / f"costs-load-{args.month}.md")
    errors = sum(i.level == ERROR for i in res.issues)
    warnings = sum(i.level == WARNING for i in res.issues)
    print(f"{args.month}: сотрудников {len(res.records)}, ошибок {errors}, предупреждений {warnings}. "
          f"Отчет: {args.out / f'costs-load-{args.month}.md'}, база: {args.db}")
    return 1 if args.strict and errors else 0


if __name__ == "__main__":
    sys.exit(main())
